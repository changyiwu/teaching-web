#!/usr/bin/env python
"""把教材檔（.pdf / .docx / .doc）抽成純文字，寫進 UTF-8 檔案。

為什麼是腳本而不是「叫 agent 每次自己寫」：`.doc` 那條路每個平台不同，
寫在散文裡等於每次重寫一次、每次重踩一次坑。

輸出一律寫檔不印到主控台——Windows 主控台是 cp950，中文 print 會爛掉
（本專案 AGENTS.md 已有三條同源的規則）。

平台差異只在 `.doc`（Word 97-2003 二進位格式）：

    Windows  win32com 驅動桌面版 Word 另存 .docx
    macOS    系統內建的 textutil（Apple 第一方，不必裝任何東西）
    其他     明確報錯並說明缺什麼

**刻意不用 LibreOffice。** `share-report/AGENTS.md` 驗證過 `soffice` 在
Windows 會因 socket.AF_UNIX 直接 AttributeError，而且它替換字型；
`cross-device-agent-skills/platform.md` 把「不要為了跨平台改用 LibreOffice」
列為明文規則，兩個平台都適用。
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
import tempfile
import uuid
from pathlib import Path


# ── .doc：唯一需要分平台的格式 ──────────────────────────────────────────

def _stage_locally(path: Path, workdir: Path) -> Path:
    """先把檔案複製到本機暫存再處理。

    雲端硬碟（Google Drive／OneDrive）是網路串流掛載，Word 直接開上面的檔案
    會因為延遲卡死；textutil 也會讀到半個檔。複製一份的成本遠低於卡死。
    """
    staged = workdir / f"{uuid.uuid4().hex}{path.suffix}"
    shutil.copy2(path, staged)
    return staged


def _doc_to_docx_windows(src: Path, workdir: Path) -> Path:
    try:
        import win32com.client  # noqa: PLC0415
    except ImportError:
        raise RuntimeError(
            "Reading .doc on Windows needs pywin32 and desktop Word."
        ) from None

    out = workdir / (src.stem + ".docx")
    word = None
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        word.DisplayAlerts = False
        # 12 = wdFormatXMLDocument（.docx）
        doc = word.Documents.Open(str(src), ConfirmConversions=False, ReadOnly=True)
        try:
            doc.SaveAs2(str(out), FileFormat=12)
        finally:
            doc.Close(SaveChanges=False)
    except Exception as exc:                      # noqa: BLE001
        raise RuntimeError(
            f"Word COM failed to convert .doc: {exc}. Is desktop Word installed?"
        ) from None
    finally:
        if word is not None:
            try:
                word.Quit()
            except Exception:                     # noqa: BLE001, S110
                pass
    return out


def _doc_to_text_macos(src: Path) -> str:
    """macOS：用系統內建的 textutil 直接轉純文字，不需要裝任何東西。"""
    exe = shutil.which("textutil")
    if not exe:
        raise RuntimeError(
            "textutil not found. It ships with macOS; its absence means this is "
            "not macOS, or PATH was changed."
        )
    proc = subprocess.run(
        [exe, "-convert", "txt", "-stdout", str(src)],
        capture_output=True, timeout=180,
    )
    if proc.returncode != 0:
        detail = proc.stderr.decode("utf-8", "replace").strip()[-300:]
        raise RuntimeError(f"textutil failed to convert .doc: {detail}")
    return proc.stdout.decode("utf-8", "replace")


def extract_doc(path: Path) -> str:
    with tempfile.TemporaryDirectory(prefix="material-") as tmp:
        workdir = Path(tmp)
        staged = _stage_locally(path, workdir)
        if sys.platform == "win32":
            return extract_docx(_doc_to_docx_windows(staged, workdir))
        if sys.platform == "darwin":
            return _doc_to_text_macos(staged)
        raise RuntimeError(
            f"No .doc reader on this platform ({sys.platform}). Windows needs "
            "desktop Word, macOS uses the built-in textutil. Re-save the file "
            "as .docx on a machine that has one of them."
        )


# ── 跨平台的兩種格式 ─────────────────────────────────────────────────────

def extract_docx(path: Path) -> str:
    import docx  # noqa: PLC0415

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    # 教材的定義、性質常放在表格裡，漏掉表格等於漏掉重點。
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)


def extract_pdf(path: Path) -> str:
    import fitz  # noqa: PLC0415

    with fitz.open(str(path)) as doc:
        return "\n".join(page.get_text() for page in doc)


EXTRACTORS = {".pdf": extract_pdf, ".docx": extract_docx, ".doc": extract_doc}


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="Extract material text (.pdf/.docx/.doc) to a UTF-8 file")
    p.add_argument("input", help="path to a .pdf, .docx or .doc file")
    p.add_argument("--out",
                   help="output .txt; defaults to the input path with .txt")
    p.add_argument("--force", action="store_true",
                   help="overwrite the output if it already exists")
    return p


def main(argv=None) -> int:
    args = build_parser().parse_args(argv)
    src = Path(args.input).expanduser().resolve()
    try:
        if not src.is_file():
            raise ValueError(f"Input file not found: {src}")
        suffix = src.suffix.lower()
        if suffix not in EXTRACTORS:
            raise ValueError(
                f"Unsupported format {suffix}; supported: {', '.join(sorted(EXTRACTORS))}"
            )
        out = Path(args.out).expanduser().resolve() if args.out \
            else src.with_suffix(".txt")
        if out.exists() and not args.force:
            raise ValueError(f"Output already exists: {out}. Use --force to overwrite.")

        text = EXTRACTORS[suffix](src)
        if not text.strip():
            raise ValueError(
                f"{src.name} produced no text. Scanned PDFs need OCR first "
                "(see the file-toolkit skill)."
            )
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(text, encoding="utf-8", newline="\n")
    except Exception as exc:                      # noqa: BLE001
        # Messages are ASCII on purpose: the Windows console is cp950 and
        # printing Chinese there is a documented failure in this project.
        print(f"{type(exc).__name__}: {exc}", file=sys.stderr)
        return 1

    print(f"OK chars={len(text)} -> {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
